"""Edge cases for extraction: encryption, numbering, encodings and nested markup."""

from __future__ import annotations

import io

import docx
import pytest
from pypdf import PdfReader, PdfWriter

from builders import build_pdf
from idengine.errors import ExtractionError
from idengine.ingest.extractors.docx import extract_docx
from idengine.ingest.extractors.html import extract_html
from idengine.ingest.extractors.markdown import extract_markdown
from idengine.ingest.extractors.pdf import extract_pdf
from idengine.ingest.extractors.plaintext import extract_plaintext
from idengine.ingest.model import BlockKind


def encrypt(data: bytes, user_password: str) -> bytes:
    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(data)).pages:
        writer.add_page(page)
    writer.encrypt(user_password=user_password, owner_password="owner-secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class TestEncryptedPdf:
    def test_empty_user_password_is_decrypted(self) -> None:
        data = encrypt(build_pdf([["Visible text in an encrypted file."]]), "")
        doc = extract_pdf(data)
        assert doc.text == "Visible text in an encrypted file."

    def test_password_protected_pdf_is_rejected(self) -> None:
        data = encrypt(build_pdf([["Secret text."]]), "letmein")
        with pytest.raises(ExtractionError, match="password"):
            extract_pdf(data)


class TestDocxNumbering:
    def test_numbered_paragraphs_become_list_items(self) -> None:
        document = docx.Document()
        document.add_paragraph("Steps to follow:")
        document.add_paragraph("Open the valve.", style="List Number")
        document.add_paragraph("Check the pressure.", style="List Number")
        buffer = io.BytesIO()
        document.save(buffer)
        doc = extract_docx(buffer.getvalue())
        assert [b.kind for b in doc.blocks] == [
            BlockKind.PARAGRAPH,
            BlockKind.LIST_ITEM,
            BlockKind.LIST_ITEM,
        ]

    def test_merged_table_cells_are_not_repeated(self) -> None:
        document = docx.Document()
        table = document.add_table(rows=2, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "Header"
        table.cell(0, 2).text = "Other"
        table.cell(1, 0).text = "a"
        table.cell(1, 1).text = "b"
        table.cell(1, 2).text = "c"
        buffer = io.BytesIO()
        document.save(buffer)
        doc = extract_docx(buffer.getvalue())
        assert doc.blocks[0].text == "Header | Other\na | b | c"

    def test_empty_document(self) -> None:
        buffer = io.BytesIO()
        docx.Document().save(buffer)
        with pytest.raises(ExtractionError, match="no text"):
            extract_docx(buffer.getvalue())


class TestHtmlEdgeCases:
    def test_nested_boilerplate_and_void_tags(self) -> None:
        html = (
            b"<html><body><nav><ul><li>Home</li><li><img src=x>About</li></ul></nav>"
            b"<article><h3>Notice</h3><p>Line one<br>line two</p>"
            b"<form><input name=q><p>search box</p></form>"
            b"<blockquote>Quoted wisdom.</blockquote></article></body></html>"
        )
        doc = extract_html(html)
        texts = [b.text for b in doc.blocks]
        assert "Home" not in " ".join(texts) and "search box" not in " ".join(texts)
        assert texts == ["Notice", "Line one", "line two", "Quoted wisdom."]
        assert doc.blocks[0].level == 3

    def test_meta_dates_and_open_graph_title(self) -> None:
        html = (
            b'<html><head><meta property="og:title" content="OG Title">'
            b'<meta property="article:published_time" content="2024-02-03T10:00:00Z">'
            b'<meta name="author" content="A. One; B. Two"></head>'
            b"<body><p>Body text.</p></body></html>"
        )
        doc = extract_html(html)
        assert doc.title == "OG Title"
        assert doc.created == "2024-02-03T10:00:00Z"
        assert doc.authors == ["A. One", "B. Two"]


class TestTextEncodings:
    def test_cp1252_plaintext(self) -> None:
        doc = extract_plaintext("Café menu\n\nCrème brûlée costs €8.".encode("cp1252"))
        assert doc.properties["encoding"] == "cp1252"
        assert "Crème brûlée" in doc.text

    def test_utf16_markdown_with_bom(self) -> None:
        doc = extract_markdown("# Título\n\nTexto en español.".encode("utf-16"))
        assert doc.title == "Título"
        assert doc.properties["encoding"] == "utf-16"

    def test_markdown_without_text_fails(self) -> None:
        with pytest.raises(ExtractionError):
            extract_markdown(b"---\ntitle: x\n---\n\n---\n")

    def test_markdown_language_and_authors_list(self) -> None:
        doc = extract_markdown(b"---\nlang: de\nauthors: [Anna, Ben]\n---\nHallo Welt.")
        assert doc.language_hint == "de"
        assert doc.authors == ["Anna", "Ben"]
        assert doc.title is None
