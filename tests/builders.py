"""Programmatic builders for binary test fixtures (PDF and DOCX)."""

from __future__ import annotations

import io
from collections.abc import Sequence

import docx


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(
    pages: Sequence[Sequence[str]],
    *,
    title: str | None = None,
    author: str | None = None,
    created: str | None = None,
) -> bytes:
    """Write a minimal, valid PDF with one Helvetica text stream per page.

    Each inner sequence holds the lines of one page; an empty string produces
    a blank line (a paragraph break after text extraction).
    """
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # placeholder, filled once the page tree id is known
    pages_id = add(b"")
    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    for lines in pages:
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 760 Td"]
        for line in lines:
            ops.append(f"({_pdf_escape(line)}) Tj T*" if line else "T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content_id = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                (
                    f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
                ).encode()
            )
        )
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode()
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    info_parts = []
    if title:
        info_parts.append(f"/Title ({_pdf_escape(title)})")
    if author:
        info_parts.append(f"/Author ({_pdf_escape(author)})")
    if created:
        info_parts.append(f"/CreationDate (D:{created})")
    info_id = add(("<< " + " ".join(info_parts) + " >>").encode()) if info_parts else None

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    trailer = f"<< /Size {len(objects) + 1} /Root {catalog} 0 R"
    if info_id:
        trailer += f" /Info {info_id} 0 R"
    out.write(f"trailer\n{trailer} >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def build_docx(
    blocks: Sequence[tuple[str, str]],
    *,
    title: str | None = None,
    author: str | None = None,
    table: Sequence[Sequence[str]] | None = None,
) -> bytes:
    """Build a DOCX from ``(kind, text)`` blocks.

    ``kind`` is ``title``, ``h1``..``h3``, ``bullet`` or ``p``. An optional
    table is appended at the end.
    """
    document = docx.Document()
    for kind, text in blocks:
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind.startswith("h"):
            document.add_heading(text, level=int(kind[1:]))
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    if title:
        document.core_properties.title = title
    if author:
        document.core_properties.author = author
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


INVOICE_TEXT = """ACME WIDGETS INVOICE

Invoice number: INV-2024-0042
Date: March 5, 2024

Payment Terms
The total amount due is $12,450.00 and payment is due within 30 days of the invoice date.
Late payments incur a 1.5% monthly fee. Contact billing@acme-widgets.example with questions.

Line Items
Industrial widgets, 500 units at $20.00 each. Installation services, 40 hours at $61.25 per hour.
"""

POLICY_MARKDOWN = """---
title: Remote Work Policy
author: Jane Doe, John Roe
date: 2023-11-01
tags: [hr, policy]
---
# Remote Work Policy

## Eligibility
Employees who have completed their probation period of 90 days may work remotely up to
three days per week, subject to manager approval.

## Equipment
The company provides a laptop and a monitor. Employees receive a stipend of $500.

## Security
All remote work must use the corporate VPN. Confidential documents must never be printed.
"""

GUIDE_HTML = """<!DOCTYPE html><html lang="en"><head><title>Kubernetes Deployment Guide</title>
<meta name="author" content="Platform Team"><meta name="description" content="How to deploy">
<script>var tracking = "do not index";</script></head>
<body><nav>Home | Docs | Login</nav>
<h1>Kubernetes Deployment Guide</h1>
<p>This guide explains how to deploy the service to a Kubernetes cluster using Helm charts.</p>
<h2>Scaling</h2>
<p>The horizontal pod autoscaler scales replicas between 2 and 10 based on CPU utilisation.</p>
<h2>Rollbacks</h2>
<p>Use helm rollback to restore the previous release when health checks fail.</p>
<table><tr><th>Env</th><th>Replicas</th></tr><tr><td>prod</td><td>6</td></tr></table>
<footer>Copyright Platform Team</footer></body></html>
"""
